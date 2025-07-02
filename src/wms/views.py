import csv
import io
import json
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from functools import wraps

import barcode
import openpyxl
import pandas as pd
from barcode.writer import ImageWriter
from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.mail import send_mail
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.db.models.functions import TruncDate
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse_lazy
from django.utils.decorators import method_decorator
from django.utils.timezone import now
from django.views import View
from django.views.generic import (CreateView, DeleteView, ListView,
                                  TemplateView, UpdateView)
from reportlab.pdfgen import canvas

from accounts.models import ROLE_CHOICES, Customer
from wms.forms import AddProductForm, ProductCreateForm
from wms.models import (OPERATION_CHOICES, Category, ChangeLog, Product,
                        StockOperation, StockOperationItem)
from wms.utils import (check_barcode_exists, get_categories_with_stats,
                       get_low_stock_products, get_products_stats,
                       get_recent_operations)


class IndexView(TemplateView):
    template_name = "index.html"


class RoleRequiredMixin:
    allowed_roles = []

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_authenticated or request.user.role not in self.allowed_roles:
            return HttpResponseForbidden("Нет доступа")
        return super().dispatch(request, *args, **kwargs)


@method_decorator(login_required, name="dispatch")
class DashboardView(RoleRequiredMixin, TemplateView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    template_name = "wms/dashboard.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = datetime.today().date()
        last_7_days = [today - timedelta(days=i) for i in reversed(range(7))]
        products_stats = get_products_stats()
        context["active_products_count"] = products_stats["count"] or 0
        context["total_quantity"] = products_stats["total_quantity"] or 0
        context["total_products_count"] = products_stats["count"] or 0
        context["total_stock_value"] = float(products_stats["total_value"] or 0)
        context["categories_data"] = get_categories_with_stats().values("name", "active_products", "total_quantity")
        category_stock_data = (
            Category.objects.annotate(total_quantity=Sum("product__quantity", filter=Q(product__is_active=True)))
            .filter(total_quantity__gt=0)
            .values("name", "total_quantity")
            .order_by("-total_quantity")
        )
        context["category_chart_labels"] = json.dumps([cat["name"] for cat in category_stock_data])
        context["category_chart_data"] = json.dumps([float(cat["total_quantity"]) for cat in category_stock_data])
        top_products_by_quantity = (
            Product.objects.filter(is_active=True, quantity__gt=0).order_by("-quantity").values("name", "quantity")[:5]
        )
        context["top_products_labels"] = json.dumps([prod["name"] for prod in top_products_by_quantity])
        context["top_products_data"] = json.dumps([float(prod["quantity"]) for prod in top_products_by_quantity])
        operations_by_day = (
            StockOperation.objects.filter(created_at__date__gte=last_7_days[0])
            .annotate(day=TruncDate("created_at"))
            .values("day", "operation_type")
            .annotate(count=Count("id"))
        )

        daily_counts = {
            int(OPERATION_CHOICES.RECEIPT): defaultdict(int),
            int(OPERATION_CHOICES.ISSUE): defaultdict(int),
            int(OPERATION_CHOICES.WRITE_OFF): defaultdict(int),
        }
        for entry in operations_by_day:
            daily_counts[entry["operation_type"]][entry["day"]] = entry["count"]
        date_labels = [day.strftime("%d-%m") for day in last_7_days]
        receipt_data = [daily_counts[OPERATION_CHOICES.RECEIPT][day] for day in last_7_days]
        issue_data = [daily_counts[OPERATION_CHOICES.ISSUE][day] for day in last_7_days]
        write_off_data = [daily_counts[OPERATION_CHOICES.WRITE_OFF][day] for day in last_7_days]
        context["date_labels"] = json.dumps(date_labels)
        context["receipt_data"] = json.dumps(receipt_data)
        context["issue_data"] = json.dumps(issue_data)
        context["write_off_data"] = json.dumps(write_off_data)
        recent_ops = get_recent_operations(10)
        context["recent_operations"] = [
            {
                "date": op.operation.created_at,
                "type": op.operation.operation_type,
                "product": op.product.name,
                "quantity": op.quantity,
            }
            for op in recent_ops
        ]
        context["low_stock_products"] = get_low_stock_products(10, 10)
        return context


@method_decorator(login_required, name="dispatch")
class ProductListView(RoleRequiredMixin, ListView):
    allowed_roles = [ROLE_CHOICES.MANAGER, ROLE_CHOICES.SELLER, ROLE_CHOICES.WORKER]
    model = Product
    template_name = "wms/products.html"
    context_object_name = "page_obj"
    paginate_by = 10

    def get_queryset(self):
        queryset = Product.objects.all()
        query = self.request.GET.get("q", "")
        category_id = self.request.GET.get("category")
        stock_status = self.request.GET.get("stock_status")
        price_range = self.request.GET.get("price_range")
        sort = self.request.GET.get("sort", "name")
        if query:
            queryset = queryset.filter(Q(name__icontains=query) | Q(barcode__icontains=query))
        if category_id:
            queryset = queryset.filter(category_id=category_id)
        if stock_status:
            if stock_status == "in_stock":
                queryset = queryset.filter(quantity__gt=0)
            elif stock_status == "low_stock":
                queryset = queryset.filter(quantity__lt=5, quantity__gt=0)
            elif stock_status == "out_of_stock":
                queryset = queryset.filter(quantity=0)
        if price_range:
            if price_range == "0-100":
                queryset = queryset.filter(selling_price__amount__lte=100)
            elif price_range == "100-500":
                queryset = queryset.filter(selling_price__amount__gt=100, selling_price__amount__lte=500)
            elif price_range == "500-1000":
                queryset = queryset.filter(selling_price__amount__gt=500, selling_price__amount__lte=1000)
            elif price_range == "1000+":
                queryset = queryset.filter(selling_price__amount__gt=1000)
        if sort == "name":
            queryset = queryset.order_by("name")
        elif sort == "-name":
            queryset = queryset.order_by("-name")
        elif sort == "quantity":
            queryset = queryset.order_by("quantity")
        elif sort == "-quantity":
            queryset = queryset.order_by("-quantity")
        elif sort == "selling_price":
            queryset = queryset.order_by("selling_price__amount")
        elif sort == "-selling_price":
            queryset = queryset.order_by("-selling_price__amount")
        elif sort == "created_at":
            queryset = queryset.order_by("created_at")
        elif sort == "-created_at":
            queryset = queryset.order_by("-created_at")
        else:
            queryset = queryset.order_by("name")
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        query = self.request.GET.get("q", "")
        category_id = self.request.GET.get("category")
        stock_status = self.request.GET.get("stock_status")
        price_range = self.request.GET.get("price_range")
        sort = self.request.GET.get("sort", "name")
        base_queryset = Product.objects.all()
        if query:
            base_queryset = base_queryset.filter(Q(name__icontains=query) | Q(barcode__icontains=query))
        if category_id:
            base_queryset = base_queryset.filter(category_id=category_id)
        if stock_status:
            if stock_status == "in_stock":
                base_queryset = base_queryset.filter(quantity__gt=0)
            elif stock_status == "low_stock":
                base_queryset = base_queryset.filter(quantity__lt=5, quantity__gt=0)
            elif stock_status == "out_of_stock":
                base_queryset = base_queryset.filter(quantity=0)
        if price_range:
            if price_range == "0-100":
                base_queryset = base_queryset.filter(selling_price__amount__lte=100)
            elif price_range == "100-500":
                base_queryset = base_queryset.filter(selling_price__amount__gt=100, selling_price__amount__lte=500)
            elif price_range == "500-1000":
                base_queryset = base_queryset.filter(selling_price__amount__gt=500, selling_price__amount__lte=1000)
            elif price_range == "1000+":
                base_queryset = base_queryset.filter(selling_price__amount__gt=1000)
        total_purchase = 0
        total_selling = 0
        total_quantity = 0
        for p in base_queryset:
            if p.purchase_price:
                total_purchase += p.purchase_price.amount
            if p.selling_price:
                total_selling += p.selling_price.amount
            total_quantity += p.quantity
        selected_category_name = ""
        if category_id:
            try:
                category = Category.objects.get(id=category_id)
                selected_category_name = category.name
            except Category.DoesNotExist:
                pass
        stock_status_display = ""
        if stock_status == "in_stock":
            stock_status_display = "В наявності"
        elif stock_status == "low_stock":
            stock_status_display = "Низький залишок"
        elif stock_status == "out_of_stock":
            stock_status_display = "Немає в наявності"
        price_range_display = ""
        if price_range == "0-100":
            price_range_display = "До 100 грн"
        elif price_range == "100-500":
            price_range_display = "100-500 грн"
        elif price_range == "500-1000":
            price_range_display = "500-1000 грн"
        elif price_range == "1000+":
            price_range_display = "Більше 1000 грн"
        context.update(
            {
                "query": query,
                "selected_category": category_id,
                "selected_category_name": selected_category_name,
                "stock_status": stock_status,
                "stock_status_display": stock_status_display,
                "price_range": price_range,
                "price_range_display": price_range_display,
                "sort": sort,
                "total_purchase": total_purchase,
                "total_selling": total_selling,
                "total_quantity": total_quantity,
                "categories": Category.objects.all(),
            }
        )
        return context


@method_decorator(login_required, name="dispatch")
class ReceiptView(RoleRequiredMixin, View):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    template_name = "wms/receipt_list.html"

    def get(self, request):
        if "receipt_cart" not in request.session:
            request.session["receipt_cart"] = []
        receipt_cart = request.session["receipt_cart"]
        add_form = AddProductForm()
        create_form = ProductCreateForm()
        products_in_cart = []
        for item in receipt_cart:
            try:
                product = Product.objects.get(id=item["product_id"])
                quantity = Decimal(item["quantity"])
                products_in_cart.append({"product": product, "quantity": quantity})
            except Product.DoesNotExist:
                continue
        search_query = request.GET.get("search", "")
        products = Product.objects.all()
        if search_query:
            products = products.filter(Q(name__icontains=search_query) | Q(barcode__icontains=search_query))
        context = {
            "products_in_cart": products_in_cart,
            "add_product_form": add_form,
            "create_product_form": create_form,
            "products": products,
            "search_query": search_query,
        }
        return render(request, self.template_name, context)

    def post(self, request):
        receipt_cart = request.session.get("receipt_cart", [])
        if "add_product" in request.POST:
            product_id = request.POST.get("product_id")
            quantity_str = request.POST.get("quantity", "0")
            if not product_id:
                messages.error(request, "Будь ласка, оберіть товар зі списку.")
                return redirect("wms:receipt_list")
            try:
                product = Product.objects.get(id=product_id)
                quantity = Decimal(quantity_str)
                if quantity <= 0:
                    raise InvalidOperation
            except (Product.DoesNotExist, InvalidOperation):
                messages.error(request, "Невірна кількість або товар.")
                return redirect("wms:receipt_list")

            for item in receipt_cart:
                if item["product_id"] == product.id:
                    item["quantity"] = str(Decimal(item["quantity"]) + quantity)
                    break
            else:
                receipt_cart.append({"product_id": product.id, "quantity": str(quantity)})

            request.session["receipt_cart"] = receipt_cart
            messages.success(request, f"Товар '{product.name}' додано.")
            return redirect("wms:receipt_list")

        elif "create_product" in request.POST:
            form = ProductCreateForm(request.POST)
            if form.is_valid():
                form.save()
                return redirect("wms:receipt_list")

        elif "clear_cart" in request.POST:
            request.session["receipt_cart"] = []
            messages.info(request, "Список очищено.")
            return redirect("wms:receipt_list")

        elif "save_receipt" in request.POST:
            if receipt_cart:
                with transaction.atomic():
                    operation = StockOperation.objects.create(
                        operation_type=OPERATION_CHOICES.RECEIPT,
                        created_by=request.user,
                        reason=request.POST.get("reason", ""),
                    )
                    for item in receipt_cart:
                        product = get_object_or_404(Product, id=item["product_id"])
                        quantity = Decimal(item["quantity"])
                        StockOperationItem.objects.create(operation=operation, product=product, quantity=quantity)
                        product.quantity += quantity
                        product.save()

                request.session["receipt_cart"] = []
                messages.success(request, "Прийом товарів збережено.")
                return redirect("wms:receipt_list")
            else:
                messages.warning(request, "Список порожній.")
                return redirect("wms:receipt_list")

        return self.get(request)


@method_decorator(login_required, name="dispatch")
class ProductCreateView(RoleRequiredMixin, CreateView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    model = Product
    form_class = ProductCreateForm
    template_name = "wms/product_create.html"
    success_url = reverse_lazy("wms:receipt_list")


@method_decorator(login_required, name="dispatch")
class ReportListView(RoleRequiredMixin, TemplateView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    template_name = "wms/reports/report_list.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["reports"] = [
            {"url": "wms:product_report", "title": "Звіт по товарам"},
            {"url": "wms:stock_report", "title": "Звіт по операціях на складі"},
            {"url": "wms:changelog_report", "title": "Звіт по змінах"},
        ]
        return context


@method_decorator(login_required, name="dispatch")
class ProductReportView(RoleRequiredMixin, ListView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    model = Product
    template_name = "wms/reports/product_report.html"
    context_object_name = "products"
    paginate_by = 10

    def get_queryset(self):
        queryset = super().get_queryset().order_by("-created_at")
        date_from = self.request.GET.get("date_from")
        date_to = self.request.GET.get("date_to")

        if date_from:
            queryset = queryset.filter(created_at__gte=date_from)
        if date_to:
            queryset = queryset.filter(created_at__lte=date_to)

        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["date_from"] = self.request.GET.get("date_from", "")
        context["date_to"] = self.request.GET.get("date_to", "")
        return context

    def render_to_response(self, context, **response_kwargs):
        export = self.request.GET.get("export")
        queryset = self.get_queryset()

        if export == "excel":
            return self.export_to_excel(queryset)
        elif export == "pdf":
            return self.export_to_pdf(queryset)
        return super().render_to_response(context, **response_kwargs)

    def export_to_excel(self, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"

        ws.append(["Назва", "Категорія", "Ціна", "Кількість", "Дата створення"])

        for p in queryset:
            ws.append(
                [
                    p.name,
                    p.category.name if p.category else "-",
                    str(p.selling_price),
                    float(p.quantity),
                    p.created_at.strftime("%Y-%m-%d") if p.created_at else "",
                ]
            )

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=products_report.xlsx"
        wb.save(response)
        return response

    def export_to_pdf(self, queryset):
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = "attachment; filename=products_report.pdf"

        p = canvas.Canvas(response)
        p.setFont("Helvetica", 12)
        y = 800
        p.drawString(100, y, "Звіт по товарах")
        y -= 30
        for product in queryset:
            line = (
                f"{product.name} | {product.category.name if product.category else '-'} | {product.selling_price} "
                f"| {product.quantity}"
            )
            p.drawString(50, y, line)
            y -= 20
            if y < 100:
                p.showPage()
                y = 800

        p.save()
        return response


@method_decorator(login_required, name="dispatch")
class StockOperationReportView(RoleRequiredMixin, ListView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    model = StockOperation
    template_name = "wms/reports/stock_operation_report.html"
    context_object_name = "operations"
    paginate_by = 10

    def get_queryset(self):
        queryset = super().get_queryset()
        date_from = self.request.GET.get("date_from")
        date_to = self.request.GET.get("date_to")
        operation_type = self.request.GET.get("operation_type")
        sort = self.request.GET.get("sort", "-created_at")
        if date_from:
            queryset = queryset.filter(created_at__gte=date_from)
        if date_to:
            queryset = queryset.filter(created_at__lte=date_to)
        if operation_type:
            queryset = queryset.filter(operation_type=operation_type)
        if sort == "created_at":
            queryset = queryset.order_by("created_at")
        elif sort == "-created_at":
            queryset = queryset.order_by("-created_at")
        elif sort == "operation_type":
            queryset = queryset.order_by("operation_type")
        elif sort == "-operation_type":
            queryset = queryset.order_by("-operation_type")
        elif sort == "created_by":
            queryset = queryset.order_by("created_by__first_name", "created_by__last_name")
        elif sort == "-created_by":
            queryset = queryset.order_by("-created_by__first_name", "-created_by__last_name")
        else:
            queryset = queryset.order_by("-created_at")

        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["date_from"] = self.request.GET.get("date_from", "")
        context["date_to"] = self.request.GET.get("date_to", "")
        context["operation_type"] = self.request.GET.get("operation_type", "")
        context["sort"] = self.request.GET.get("sort", "-created_at")
        context["operation_choices"] = StockOperation._meta.get_field("operation_type").choices
        return context

    def render_to_response(self, context, **response_kwargs):
        export = self.request.GET.get("export")
        queryset = self.get_queryset()

        if export == "excel":
            return self.export_to_excel(queryset)
        elif export == "pdf":
            return self.export_to_pdf(queryset)

        return super().render_to_response(context, **response_kwargs)

    def export_to_excel(self, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Stock Operations"

        ws.append(["Тип операції", "Користувач", "Причина", "Примітка", "Товари", "Дата"])

        for op in queryset:
            items_str = "; ".join([f"{item.product.name} ({item.quantity})" for item in op.items.all()])
            ws.append(
                [
                    op.get_operation_type_display(),
                    op.created_by.get_full_name() if op.created_by else "-",
                    op.reason or "-",
                    op.note or "-",
                    items_str,
                    op.created_at.strftime("%Y-%m-%d %H:%M"),
                ]
            )

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=stock_operations_report.xlsx"
        wb.save(response)
        return response

    def export_to_pdf(self, queryset):
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = "attachment; filename=stock_operations_report.pdf"

        p = canvas.Canvas(response)
        p.setFont("Helvetica", 12)
        y = 800
        p.drawString(50, y, "Звіт по операціях на складі")
        y -= 30

        for op in queryset:
            items_str = ", ".join([f"{item.product.name}({item.quantity})" for item in op.items.all()])
            line = (
                f"{op.get_operation_type_display()} | {op.created_by.get_full_name() if op.created_by else '-'} |"
                f" {op.reason or '-'} | {items_str} | {op.created_at.strftime('%Y-%m-%d %H:%M')}"
            )
            p.drawString(20, y, line)
            y -= 20
            if y < 100:
                p.showPage()
                y = 800

        p.save()
        return response


@method_decorator(login_required, name="dispatch")
class ChangeLogReportView(RoleRequiredMixin, TemplateView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    template_name = "wms/reports/changelog_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        last_30_days = now() - timedelta(days=30)
        changelogs = ChangeLog.objects.filter(created_at__gte=last_30_days).order_by("-created_at")
        context["changelogs"] = changelogs
        return context


@method_decorator(login_required, name="dispatch")
class IssueView(RoleRequiredMixin, View):
    allowed_roles = [ROLE_CHOICES.MANAGER, ROLE_CHOICES.SELLER]
    template_name = "wms/issue.html"

    def get(self, request):
        issue_cart = request.session.get("issue_cart", [])
        total_items = sum(Decimal(item["quantity"]) for item in issue_cart)

        context = {
            "issue_cart": issue_cart,
            "total_items": total_items,
            "products": Product.objects.filter(is_active=True).order_by("name"),
        }
        return render(request, self.template_name, context)

    def post(self, request):
        issue_cart = request.session.get("issue_cart", [])

        if "add_product" in request.POST:
            product_id = request.POST.get("product")
            try:
                quantity = Decimal(request.POST.get("quantity", 0))
            except InvalidOperation:
                messages.error(request, "Некоректна кількість")
                return redirect("wms:issue_list")

            if quantity <= 0:
                messages.error(request, "Кількість має бути більшою за 0")
                return redirect("wms:issue_list")

            product = get_object_or_404(Product, id=product_id)

            if product.quantity < quantity:
                messages.error(request, f"Недостатньо товару на складі. Доступно: {product.quantity}")
                return redirect("wms:issue_list")

            for item in issue_cart:
                if item["product_id"] == int(product_id):
                    new_quantity = Decimal(item["quantity"]) + quantity
                    if product.quantity < new_quantity:
                        messages.error(request, f"Недостатньо товару на складі. Доступно: {product.quantity}")
                        return redirect("wms:issue_list")
                    item["quantity"] = str(new_quantity)
                    break
            else:
                issue_cart.append(
                    {
                        "product_id": int(product_id),
                        "product_name": product.name,
                        "quantity": str(quantity),
                        "unit": product.get_unit_display(),
                    }
                )

            request.session["issue_cart"] = issue_cart
            messages.success(request, "Товар додано до списку видачі")
            return redirect("wms:issue_list")

        elif "remove_product" in request.POST:
            product_id = request.POST.get("product_id")
            issue_cart = [item for item in issue_cart if item["product_id"] != int(product_id)]
            request.session["issue_cart"] = issue_cart
            messages.success(request, "Товар видалено зі списку")
            return redirect("wms:issue_list")

        elif "save_issue" in request.POST:
            if issue_cart:
                with transaction.atomic():
                    operation = StockOperation.objects.create(
                        operation_type=OPERATION_CHOICES.ISSUE,
                        created_by=request.user,
                        reason=request.POST.get("reason", ""),
                        note=request.POST.get("note", ""),
                    )

                    for item in issue_cart:
                        product = get_object_or_404(Product, id=item["product_id"])
                        quantity = Decimal(item["quantity"])

                        if product.quantity < quantity:
                            messages.error(request, f"Недостатньо товару на складі. Доступно: {product.quantity}")
                            return redirect("wms:issue_list")

                        StockOperationItem.objects.create(operation=operation, product=product, quantity=quantity)
                        product.quantity -= quantity
                        product.save()

                request.session["issue_cart"] = []
                messages.success(request, "Видачу товарів збережено")
                return redirect("wms:issue_list")
            else:
                messages.warning(request, "Список порожній")
                return redirect("wms:issue_list")

        return self.get(request)


@method_decorator(login_required, name="dispatch")
class ProductUpdateView(RoleRequiredMixin, UpdateView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    model = Product
    form_class = ProductCreateForm
    template_name = "wms/product_form.html"
    success_url = reverse_lazy("wms:product_list")


@method_decorator(login_required, name="dispatch")
class ProductDeleteView(RoleRequiredMixin, DeleteView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    model = Product
    template_name = "wms/product_confirm_delete.html"
    success_url = reverse_lazy("wms:product_list")

    def delete(self, request, *args, **kwargs):
        product = self.get_object()
        messages.success(request, f"Товар '{product.name}' успішно видалено.")
        return super().delete(request, *args, **kwargs)


@method_decorator(login_required, name="dispatch")
class WriteOffView(RoleRequiredMixin, View):
    allowed_roles = [ROLE_CHOICES.MANAGER, ROLE_CHOICES.SELLER]
    template_name = "wms/writeoff.html"

    def get(self, request):
        writeoff_cart = request.session.get("writeoff_cart", [])
        total_items = sum(Decimal(item["quantity"]) for item in writeoff_cart)

        context = {
            "writeoff_cart": writeoff_cart,
            "total_items": total_items,
            "products": Product.objects.filter(is_active=True).order_by("name"),
        }
        return render(request, self.template_name, context)

    def post(self, request):
        writeoff_cart = request.session.get("writeoff_cart", [])

        if "add_product" in request.POST:
            product_id = request.POST.get("product")
            try:
                quantity = Decimal(request.POST.get("quantity", 0))
            except InvalidOperation:
                messages.error(request, "Неверное количество")
                return redirect("wms:writeoff_list")

            if quantity <= 0:
                messages.error(request, "Кількість має бути більшою за 0")
                return redirect("wms:writeoff_list")

            product = get_object_or_404(Product, id=product_id)

            if product.quantity < quantity:
                messages.error(request, f"Недостатньо товару на складі. Доступно: {product.quantity}")
                return redirect("wms:writeoff_list")

            for item in writeoff_cart:
                if item["product_id"] == int(product_id):
                    new_quantity = Decimal(item["quantity"]) + quantity
                    if product.quantity < new_quantity:
                        messages.error(request, f"Недостатньо товару на складі. Доступно: {product.quantity}")
                        return redirect("wms:writeoff_list")
                    item["quantity"] = str(new_quantity)
                    break
            else:
                writeoff_cart.append(
                    {
                        "product_id": int(product_id),
                        "product_name": product.name,
                        "quantity": str(quantity),
                        "unit": product.get_unit_display(),
                    }
                )

            request.session["writeoff_cart"] = writeoff_cart
            messages.success(request, "Товар додано до списку списання")
            return redirect("wms:writeoff_list")

        elif "remove_product" in request.POST:
            product_id = request.POST.get("product_id")
            writeoff_cart = [item for item in writeoff_cart if item["product_id"] != int(product_id)]
            request.session["writeoff_cart"] = writeoff_cart
            messages.success(request, "Товар видалено зі списку")
            return redirect("wms:writeoff_list")

        elif "save_writeoff" in request.POST:
            if writeoff_cart:
                with transaction.atomic():
                    operation = StockOperation.objects.create(
                        operation_type=OPERATION_CHOICES.WRITE_OFF,
                        created_by=request.user,
                        reason=request.POST.get("reason", ""),
                        note=request.POST.get("note", ""),
                    )

                    for item in writeoff_cart:
                        product = get_object_or_404(Product, id=item["product_id"])
                        quantity = Decimal(item["quantity"])

                        if product.quantity < quantity:
                            messages.error(request, f"Недостатньо товару на складі. Доступно: {product.quantity}")
                            return redirect("wms:writeoff_list")

                        StockOperationItem.objects.create(operation=operation, product=product, quantity=quantity)
                        product.quantity -= quantity
                        product.save()

                request.session["writeoff_cart"] = []
                messages.success(request, "Списання товарів збережено")
                return redirect("wms:writeoff_list")
            else:
                messages.warning(request, "Список порожній")
                return redirect("wms:writeoff_list")

        return self.get(request)


@method_decorator(login_required, name="dispatch")
class ToolsView(RoleRequiredMixin, TemplateView):
    allowed_roles = [ROLE_CHOICES.MANAGER]
    template_name = "wms/tools.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        import random

        while True:
            code = str(random.randint(10**11, 10**12 - 1))
            if not check_barcode_exists(code):
                context["default_barcode"] = code
                break
        return context


@login_required
def barcode_generator(request):
    if request.method == "GET" and request.GET.get("free") == "1":
        import random

        while True:
            code = str(random.randint(10**11, 10**12 - 1))  # 12-значный
            if not check_barcode_exists(code):
                return JsonResponse({"barcode": code})
    if request.method == "POST":
        import json as pyjson

        try:
            data = pyjson.loads(request.body.decode())
            code = data.get("text", "")
        except Exception:
            code = request.POST.get("barcode_text", "")
        if not code:
            return JsonResponse({"success": False, "error": "Введіть текст для генерації штрихкоду"}, status=400)
        if check_barcode_exists(code):
            return JsonResponse(
                {"success": False, "error": "Такий штрихкод вже існує в базі! Введіть інший."}, status=400
            )
        ean = barcode.get("ean13", code.zfill(12), writer=ImageWriter())
        buffer = io.BytesIO()
        ean.write(buffer)
        buffer.seek(0)
        import base64

        img_base64 = base64.b64encode(buffer.read()).decode("utf-8")
        img_url = f"data:image/png;base64,{img_base64}"
        return JsonResponse({"success": True, "barcode_url": img_url})
    return HttpResponseForbidden()


def manager_required(view_func):

    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect("login")
        if request.user.role != ROLE_CHOICES.MANAGER:
            messages.error(request, "Доступ заборонено. Потрібні права менеджера.")
            return redirect("wms:dashboard")
        return view_func(request, *args, **kwargs)

    return _wrapped_view


def manager_required_mixin(view_class):

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect("login")
        if request.user.role != ROLE_CHOICES.MANAGER:
            messages.error(request, "Доступ заборонено. Потрібні права менеджера.")
            return redirect("wms:dashboard")
        return super().dispatch(request, *args, **kwargs)

    view_class.dispatch = dispatch
    return view_class


@manager_required
def export_products_csv(request):
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="products_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv"'

    writer = csv.writer(response)
    writer.writerow(
        [
            "name",
            "barcode",
            "category",
            "quantity",
            "unit",
            "purchase_price",
            "sale_price",
            "description",
        ]
    )

    products = Product.objects.filter(is_active=True)
    for product in products:
        writer.writerow(
            [
                product.name,
                product.barcode or "",
                product.category.name if product.category else "",
                product.quantity,
                product.unit,
                product.purchase_price.amount if product.purchase_price else "",
                product.selling_price.amount if product.selling_price else "",
                product.description or "",
            ]
        )

    return response


@manager_required
def export_products_excel(request):
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="products_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Товари"

    headers = [
        "Назва",
        "Штрихкод",
        "Категорія",
        "Кількість",
        "Одиниця",
        "Закупівельна ціна",
        "Продажна ціна",
        "Опис",
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    products = Product.objects.filter(is_active=True)
    for row, product in enumerate(products, 2):
        ws.cell(row=row, column=1, value=product.name)
        ws.cell(row=row, column=2, value=product.barcode or "")
        ws.cell(row=row, column=3, value=product.category.name if product.category else "")
        ws.cell(row=row, column=4, value=float(product.quantity))
        ws.cell(row=row, column=5, value=product.get_unit_display())
        ws.cell(row=row, column=6, value=float(product.purchase_price.amount) if product.purchase_price else "")
        ws.cell(row=row, column=7, value=float(product.selling_price.amount) if product.selling_price else "")
        ws.cell(row=row, column=8, value=product.description or "")

    wb.save(response)
    return response


class ImportProductsForm(forms.Form):
    file = forms.FileField(label="Оберіть файл (CSV або Excel)")


@manager_required
def import_products(request):
    print("import_products: start")
    if request.method == "POST":
        print("import_products: POST")
        file = request.FILES.get("file")
        if not file:
            print("import_products: no file")
            messages.error(request, "Будь ласка, виберіть файл для імпорту.")
            form = ImportProductsForm()
            return render(request, "wms/import_products.html", {"form": form})
        try:
            print(f"import_products: file name = {file.name}")
            if file.name.endswith(".csv"):
                df = pd.read_csv(file)
            elif file.name.endswith((".xlsx", ".xls")):
                df = pd.read_excel(file)
            else:
                print("import_products: unsupported file format")
                messages.error(
                    request,
                    "Непідтримуваний формат файлу. Використовуйте CSV або Excel.",
                )
                form = ImportProductsForm()
                return render(request, "wms/import_products.html", {"form": form})
            required_columns = ["name", "quantity", "purchase_price"]
            has_sale_price = "sale_price" in df.columns
            has_selling_price = "selling_price" in df.columns
            if not (has_sale_price or has_selling_price):
                missing_columns = ["sale_price (або selling_price)"]
            else:
                missing_columns = []
            for col in required_columns:
                if col not in df.columns:
                    missing_columns.append(col)
            print(f"import_products: missing_columns = {missing_columns}")
            if missing_columns:
                messages.error(request, f'Відсутні обов\'язкові колонки: {", ".join(missing_columns)}')
                form = ImportProductsForm()
                return render(request, "wms/import_products.html", {"form": form})
            if has_selling_price and not has_sale_price:
                df["sale_price"] = df["selling_price"]
            request.session["import_data"] = df.to_dict("records")
            request.session["import_filename"] = file.name
            print("import_products: success, redirect")
            return redirect("wms:preview_import_products")
        except Exception as e:
            import traceback

            tb = traceback.format_exc()
            print(f"import_products: exception: {e}\n{tb}")
            messages.error(request, f"Помилка при обробці файлу: {str(e)}\n{tb}")
            form = ImportProductsForm()
            return render(request, "wms/import_products.html", {"form": form})
    else:
        print("import_products: GET")
        form = ImportProductsForm()
    print("import_products: render form")
    return render(request, "wms/import_products.html", {"form": form})


@manager_required
def preview_import_products(request):
    import_data = request.session.get("import_data")
    filename = request.session.get("import_filename")

    if not import_data:
        messages.error(request, "Немає даних для імпорту.")
        return redirect("wms:tools")

    preview_data = import_data[:5]

    if request.method == "POST":
        try:
            created_count = 0
            updated_count = 0

            for row in import_data:
                name = row.get("name", "").strip()
                if not name:
                    continue

                product, created = Product.objects.get_or_create(
                    name=name,
                    defaults={
                        "quantity": float(row.get("quantity", 0)),
                        "purchase_price": float(row.get("purchase_price", 0)),
                        "selling_price": float(row.get("sale_price", 0)),
                        "barcode": row.get("barcode", "").strip(),
                        "description": row.get("description", "").strip(),
                        "unit": int(row.get("unit", 1)),
                    },
                )

                if created:
                    created_count += 1
                else:
                    product.quantity = float(row.get("quantity", 0))
                    product.purchase_price = float(row.get("purchase_price", 0))
                    product.selling_price = float(row.get("sale_price", 0))
                    product.save()
                    updated_count += 1

            del request.session["import_data"]
            del request.session["import_filename"]

            messages.success(
                request,
                f"Імпорт завершено! Створено: {created_count}, оновлено: {updated_count}",
            )
            return redirect("wms:import_products")

        except Exception as e:
            messages.error(request, f"Помилка при імпорті: {str(e)}")
            return redirect("wms:tools")

    context = {"preview_data": preview_data, "filename": filename, "total_rows": len(import_data)}
    return render(request, "wms/import_products.html", context)


@manager_required
def backup_products(request):
    from django.core.serializers import serialize

    products = Product.objects.filter(is_active=True)
    data = serialize("json", products)

    response = HttpResponse(data, content_type="application/json")
    response["Content-Disposition"] = (
        f'attachment; filename="products_backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.json"'
    )

    messages.success(request, "Резервна копія товарів створена успішно!")
    return response


@manager_required
def restore_products(request):
    if request.method == "POST":
        file = request.FILES.get("backup_file")
        if not file:
            messages.error(request, "Будь ласка, виберіть файл резервної копії.")
            return redirect("wms:tools")

        try:
            data = json.load(file)
            restored_count = 0

            for item in data:
                if item["model"] == "wms.product":
                    fields = item["fields"]
                    product, created = Product.objects.get_or_create(
                        name=fields["name"],
                        defaults={
                            "barcode": fields.get("barcode", ""),
                            "purchase_price": fields.get("purchase_price", 0),
                            "selling_price": fields.get("selling_price", 0),
                            "unit": fields.get("unit", 1),
                            "quantity": fields.get("quantity", 0),
                            "is_active": fields.get("is_active", True),
                            "description": fields.get("description", ""),
                        },
                    )
                    if created:
                        restored_count += 1

            messages.success(request, f"Відновлено {restored_count} товарів!")
            return redirect("wms:product_list")

        except Exception as e:
            messages.error(request, f"Помилка при відновленні: {str(e)}")
            return redirect("wms:tools")

    return render(request, "wms/restore_products.html")


@manager_required
def backup_all_data(request):
    from django.core.serializers import serialize

    from accounts.models import Customer

    data = []

    users = Customer.objects.all()
    data.extend(serialize("python", users))

    categories = Category.objects.all()
    data.extend(serialize("python", categories))

    products = Product.objects.all()
    data.extend(serialize("python", products))

    operations = StockOperation.objects.all()
    data.extend(serialize("python", operations))

    operation_items = StockOperationItem.objects.all()
    data.extend(serialize("python", operation_items))

    changes = ChangeLog.objects.all()
    data.extend(serialize("python", changes))

    response = HttpResponse(json.dumps(data, indent=2), content_type="application/json")
    response["Content-Disposition"] = (
        f'attachment; filename="full_backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.json"'
    )

    messages.success(request, "Повна резервна копія створена успішно!")
    return response


@manager_required
def restore_all_data(request):
    if request.method == "POST":
        file = request.FILES.get("backup_file")
        if not file:
            messages.error(request, "Будь ласка, виберіть файл резервної копії.")
            return redirect("wms:tools")

        try:
            data = json.load(file)
            restored_counts = {
                "users": 0,
                "categories": 0,
                "products": 0,
                "operations": 0,
                "operation_items": 0,
                "changes": 0,
            }

            for item in data:
                model_name = item["model"]
                fields = item["fields"]

                if model_name == "accounts.customer":
                    user, created = Customer.objects.get_or_create(email=fields["email"], defaults=fields)
                    if created:
                        restored_counts["users"] += 1

                elif model_name == "wms.category":
                    category, created = Category.objects.get_or_create(name=fields["name"], defaults=fields)
                    if created:
                        restored_counts["categories"] += 1

                elif model_name == "wms.product":
                    product, created = Product.objects.get_or_create(name=fields["name"], defaults=fields)
                    if created:
                        restored_counts["products"] += 1

                elif model_name == "wms.stockoperation":
                    operation, created = StockOperation.objects.get_or_create(id=item["pk"], defaults=fields)
                    if created:
                        restored_counts["operations"] += 1

                elif model_name == "wms.stockoperationitem":
                    item_obj, created = StockOperationItem.objects.get_or_create(id=item["pk"], defaults=fields)
                    if created:
                        restored_counts["operation_items"] += 1

                elif model_name == "wms.changelog":
                    change, created = ChangeLog.objects.get_or_create(id=item["pk"], defaults=fields)
                    if created:
                        restored_counts["changes"] += 1

            total_restored = sum(restored_counts.values())
            messages.success(request, f"Відновлено {total_restored} записів!")
            return redirect("wms:dashboard")

        except Exception as e:
            messages.error(request, f"Помилка при відновленні: {str(e)}")
            return redirect("wms:tools")

    return render(request, "wms/restore_all_data.html")


class HelpView(TemplateView):
    template_name = "wms/help/help_main.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["help_sections"] = [
            {
                "title": "Головна",
                "icon": "bi-house",
                "description": "Загальний огляд системи WMS",
                "url": "wms:help_main",
            },
            {
                "title": "Товари",
                "icon": "bi-box-seam",
                "description": "Управління товарами та категоріями",
                "url": "wms:help_products",
            },
            {
                "title": "Операції",
                "icon": "bi-arrow-left-right",
                "description": "Прийом, видача та списання товарів",
                "url": "wms:help_operations",
            },
            {
                "title": "Звіти",
                "icon": "bi-clipboard-data",
                "description": "Генерація та перегляд звітів",
                "url": "wms:help_reports",
            },
            {
                "title": "Інструменти",
                "icon": "bi-tools",
                "description": "Імпорт/експорт та резервне копіювання",
                "url": "wms:help_tools",
            },
            {
                "title": "FAQ",
                "icon": "bi-question-circle",
                "description": "Часто задавані питання",
                "url": "wms:help_faq",
            },
        ]
        return context


class HelpProductsView(TemplateView):
    template_name = "wms/help/help_products.html"


class HelpOperationsView(TemplateView):
    template_name = "wms/help/help_operations.html"


class HelpReportsView(TemplateView):
    template_name = "wms/help/help_reports.html"


class HelpToolsView(TemplateView):
    template_name = "wms/help/help_tools.html"


class HelpFaqView(TemplateView):
    template_name = "wms/help/help_faq.html"


class BarcodeGeneratorView(LoginRequiredMixin, TemplateView):
    template_name = "wms/barcode_generator.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        import random

        while True:
            code = str(random.randint(10**11, 10**12 - 1))
            if not Product.objects.filter(barcode=code).exists():
                context["default_barcode"] = code
                break
        return context


@manager_required
def preview_import_products_ajax(request):
    if request.method == "POST" and request.headers.get("x-requested-with") == "XMLHttpRequest":
        file = request.FILES.get("file")
        if not file:
            print("preview_import_products_ajax: no file")
            return JsonResponse({"error": "Будь ласка, виберіть файл для імпорту."})
        try:
            print(f"preview_import_products_ajax: file name = {file.name}")
            if file.name.endswith(".csv"):
                df = pd.read_csv(file)
            elif file.name.endswith((".xlsx", ".xls")):
                df = pd.read_excel(file)
            else:
                print("preview_import_products_ajax: unsupported file format")
                return JsonResponse({"error": "Непідтримуваний формат файлу. Використовуйте CSV або Excel."})
            print(f"preview_import_products_ajax: columns = {list(df.columns)}")
            preview_data = df.head(5).to_dict("records")
            html = render_to_string("wms/import_preview_table.html", {"preview_data": preview_data})
            print("preview_import_products_ajax: success")
            return JsonResponse({"html": html})
        except Exception as e:
            import traceback

            tb = traceback.format_exc()
            print(f"preview_import_products_ajax: exception: {e}\n{tb}")
            return JsonResponse({"error": f"Помилка при обробці файлу: {str(e)}\n{tb}"})
    print("preview_import_products_ajax: некоректний запит")
    return JsonResponse({"error": "Некоректний запит."})


@manager_required
def send_email_to_managers(request):
    if request.method == "POST":
        recipient_type = request.POST.get("recipient_type")
        custom_email = request.POST.get("custom_email", "").strip()
        subject = request.POST.get("subject", "").strip()
        message = request.POST.get("message", "").strip()

        if not subject or not message:
            messages.error(request, "Тема та повідомлення обов'язкові.")
            return redirect("wms:tools")

        recipients = []

        if recipient_type == "managers":
            managers = Customer.objects.filter(role=ROLE_CHOICES.MANAGER, is_active=True)
            recipients = [manager.email for manager in managers]
        elif recipient_type == "employees":
            employees = Customer.objects.filter(
                Q(role=ROLE_CHOICES.SELLER) | Q(role=ROLE_CHOICES.WORKER), is_active=True
            )
            recipients = [employee.email for employee in employees]
        elif recipient_type == "custom" and custom_email:
            recipients = [email.strip() for email in custom_email.split(",") if email.strip()]

        if not recipients:
            messages.error(request, "Не вказано отримувачів email.")
            return redirect("wms:tools")

        try:
            send_mail(
                subject=subject,
                message=message,
                from_email=None,  # Використовуємо DEFAULT_FROM_EMAIL з налаштувань
                recipient_list=recipients,
                fail_silently=False,
            )

            messages.success(request, f"Email успішно відправлено {len(recipients)} отримувачам.")

        except Exception as e:
            messages.error(request, f"Помилка при відправці email: {str(e)}")

        return redirect("wms:tools")

    context = {
        "managers_count": Customer.objects.filter(role=ROLE_CHOICES.MANAGER, is_active=True).count(),
        "employees_count": Customer.objects.filter(
            Q(role=ROLE_CHOICES.SELLER) | Q(role=ROLE_CHOICES.WORKER), is_active=True
        ).count(),
    }

    return render(request, "wms/send_email.html", context)
